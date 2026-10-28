from pathlib import Path
import json
from decimal import Decimal, ROUND_HALF_UP
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'documents/milestones'; QA=OUT/'source'; FIG=OUT/'figures'
E=json.loads((QA/'evidence.json').read_text()); M={k:v['metrics'] for k,v in E['subjects'].items()}
def newdoc():
 d=Document(); s=d.sections[0]; s.page_height=Inches(11);s.page_width=Inches(8.5)
 s.top_margin=s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.76)
 for name,size in [('Normal',11),('Title',27),('Subtitle',14),('Heading 1',18),('Heading 2',12)]:
  z=d.styles[name];z.font.name='Arial';z.font.size=Pt(size);z.font.color.rgb=RGBColor(0,0,0)
  z.paragraph_format.space_after=Pt(7);z.paragraph_format.line_spacing=1.1
 d.styles['Heading 2'].paragraph_format.space_before=Pt(8)
 for border in d.styles.element.xpath('.//w:pBdr'):
  border.getparent().remove(border)
 header=s.header.paragraphs[0];header.text='StudentInsight   |   Bindura University of Science Education';header.runs[0].font.size=Pt(8)
 foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 foot.add_run('StudentInsight   |   ').font.size=Pt(8)
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
 return d
def p(text):
 for chunk in text.split('\n'):
  if chunk:doc.add_paragraph(chunk)
def table(headers, rows, widths=None):
 widths=widths or ([2,3.3,1.45] if len(headers)==3 else [6.75/len(headers)]*len(headers))
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,x in zip(t.rows[0].cells,headers):c.text=str(x)
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for c,w in zip(row.cells,widths):
   c.width=Inches(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'183844' if i==0 else ('F1F5F6' if i%2 else 'FFFFFF'));cp.append(sh)
   mar=OxmlElement('w:tcMar');borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');mar.append(el)
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   cp.append(mar);cp.append(borders)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.line_spacing=1.05
    for r in pp.runs:r.font.size=Pt(9);r.font.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
def diagram():
 im=Image.new('RGB',(1500,670),'white');dr=ImageDraw.Draw(im)
 f=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',27);b=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',30)
 def box(x,y,title,lines):
  dr.rectangle((x,y,x+435,y+185),fill='#F1F5F6',outline='#183844',width=3);dr.text((x+18,y+15),title,font=b,fill='#183844')
  for j,line in enumerate(lines):dr.text((x+18,y+62+j*33),line,font=f,fill='black')
 def arrow(x,y,xx,yy):
  dr.line((x,y,xx,yy),fill='#087F8C',width=5);dr.polygon([(xx,yy),(xx-12,yy-10),(xx-12,yy+10)],fill='#087F8C')
 box(20,25,'Staff interface',['Sign in and subject selection','Individual and class inputs','auth.py and app.py'])
 box(530,25,'Validation and prediction',['Active feature schema','kaggle_model.py','Performance band estimates'])
 box(1040,25,'Visible outputs',['Charts and review tables','Bands and model quality','CSV reports'])
 arrow(455,115,530,115);arrow(965,115,1040,115)
 box(20,315,'Recorded data',['Kaggle xAPI and UCI reference','Private class data in session','Metadata outside predictors'])
 box(530,315,'Training and evaluation',['Grouped validation folds','Select then fit then test','Recorded evaluation evidence'])
 box(1040,315,'Notebook integration',['Jupyter and Google Colab','Optional Drive access','Selected CSV and ZIP reports'])
 arrow(455,405,530,405);arrow(965,405,1040,405)
 dr.line((745,315,745,210),fill='#087F8C',width=5);dr.polygon([(745,210),(735,225),(755,225)],fill='#087F8C')
 dr.text((30,585),'Public bundles use shared caches. Private models and records stay in the staff session.',font=f,fill='#183844')
 im.save(FIG/'architecture.png')
diagram();doc=newdoc()
doc.add_paragraph('StudentInsight',style='Title');doc.add_paragraph('Student Performance Prediction',style='Title')
doc.add_paragraph('Technical report covering five milestone chapters',style='Subtitle')
for line in ['Bindura University of Science Education','Programme  Network Engineering','Module  Artificial Intelligence and Intelligent System','Supervisor  Mr C. Zano','Submission  Group project','Group number  ____________________','Group members and student numbers']:
 p(line)
for _ in range(4):p('________________________________________________________')
p('Updated 28 October 2026 | Application baseline 5af3632')
p('This report follows the supplied M1 to M5 structure. It documents the implemented system and measured results. Complete group identities and role allocations before submission.')
for entry in json.loads((QA/'report_content.json').read_text(encoding='utf-8')):
 doc.add_page_break();doc.add_heading(entry['title'],1)
 for text in entry.get('paragraphs',[]):p(text)
 for section in entry.get('sections',[]):doc.add_heading(section['heading'],2);p(section['text'])
 for text in entry.get('bullets',[]):doc.add_paragraph(text,style='List Bullet')
 if 'figure' in entry:
  doc.add_picture(str(FIG/entry['figure']),width=Inches(6.7));cap=doc.add_paragraph(entry['caption']);cap.runs[0].italic=True;cap.runs[0].font.size=Pt(9)
 if 'paired_figures' in entry:
  # Lay out unaltered screen captures on one canvas to avoid inline-image wrapping.
  screens=[Image.open(FIG/filename).convert('RGB') for filename in entry['paired_figures']]
  canvas=Image.new('RGB',(sum(im.width for im in screens)+30,max(im.height for im in screens)),'white')
  x=0
  for screen in screens:
   canvas.paste(screen,(x,0));x+=screen.width+30
  canvas.save(FIG/'account-navigation.png')
  pp=doc.add_paragraph();pp.paragraph_format.keep_with_next=True
  pp.add_run().add_picture(str(FIG/'account-navigation.png'),width=Inches(6.35))
  cap=doc.add_paragraph(entry['pair_caption']);cap.runs[0].italic=True;cap.runs[0].font.size=Pt(9)
 for figure in entry.get('figures',[]):
  pp=doc.add_paragraph();pp.paragraph_format.keep_with_next=True
  pp.add_run().add_picture(str(FIG/figure['file']),width=Inches(figure.get('width',6.7)))
  cap=doc.add_paragraph(figure['caption']);cap.runs[0].italic=True;cap.runs[0].font.size=Pt(9)
 if 'table' in entry:
  t=entry['table'];table(t['headers'],t['rows'],[.65,1.85,4.25] if t['headers'][0]=='Chapter' else None)
 dynamic=entry.get('dynamic')
 if dynamic=='cv':
  rows=[]
  for i,c in enumerate(M['mat']['comparison']):
   q=M['por']['comparison'][i];rows.append([c['model'],c['feature_set'],f"{c['cv_mae']:.3f} / {c['cv_std']:.3f}",f"{q['cv_mae']:.3f} / {q['cv_std']:.3f}"])
  table(['Candidate','Schema','Math MAE / SD','Portuguese MAE / SD'],rows,[2.25,1.35,1.5,1.65])
 if dynamic=='results':
  rows=[]
  for key in ['mae','rmse','r2']:
   rows.append([key.upper(),*[f"{E['legacy']['subjects'][s]['reproduced_metrics'][key]:.3f} to {M[s][key]:.3f}" for s in ['mat','por']]])
  rows += [['Interval coverage','68.4% to 83.5%','79.2% to 93.8%'],['Average width','2.896 to 2.000','2.013 to 2.000']]
  table(['Measure','Mathematics before to after','Portuguese before to after'],rows,[1.65,2.55,2.55])
 if dynamic=='alerts':
  rows=[]
  for s in ['mat','por']:
   a=E['subjects'][s]['alerts'];rows.append([M[s]['subject_label'],a['tp'],a['fp'],a['fn'],a['tn'],str(Decimal(str(a['precision'])).quantize(Decimal('.001'),rounding=ROUND_HALF_UP)),f"{a['recall']:.3f}",f"{a['f1']:.3f}"])
  table(['Subject','TP','FP','FN','TN','Precision','Recall','F1'],rows,[1.5,.5,.5,.5,.5,1.1,1.1,1.05])
 if dynamic=='tests':table(['Test module','Passing tests'],list(E['tests'].items())+[['Total',E['test_count']]],[5.35,1.4])
 if dynamic=='milestones':
  doc.add_heading('5 3 Assessment rubric coverage',2)
  table(['Criterion','Marks','Evidence'],[['Problem definition and design',3,'Chapters 1 and 2'],['Technical implementation',4,'Chapters 3 and 4; source and code images'],['Evaluation and analysis',3,'Chapter 3 metrics and limitations'],['Documentation and GitHub',2,'Chapter 4; README; versioned repository'],['Presentation and demo',2,'Slides; 15-minute guide; delivery pending'],['Ethical consideration',1,'Chapter 2; privacy and model limitations']],[2.7,.55,3.5])
  p('Total: 15 marks, representing 15% of the course grade in the supplied brief. This table maps evidence to criteria; it does not award marks. The technical report limit is 20 pages, including this edition’s cover and references.')
 if 'after' in entry:p(entry['after'])
doc.save(OUT/'StudentInsight_Five_Chapter_Report.docx')

doc=newdoc();doc.add_paragraph('StudentInsight demonstration guide',style='Title')
p('Bindura University of Science Education\nNetwork Engineering\nArtificial Intelligence and Intelligent System\nSupervisor Mr C. Zano')
p('Duration 15 minutes. This guide is separate from the technical report. Assign speakers and an operator, then rehearse the sequence before assessment.')
doc.add_heading('Preparation',2)
for x in ['Complete the group names, student numbers and roles on the submission materials.', 'Start Streamlit using .venv on port 8502 and sign in with an authorized demonstration account. Use a demonstration account and keep passwords out of captures.', 'Open the slides, notebook and Model quality view. Use bundled historical records first. Clearly label any synthetic class fixtures.', 'Authorize Drive before presenting if showing the cloud step. Otherwise explain that code exists but cloud verification is pending.', 'Use Kaggle academic performance by default. Explain Low, Middle and High bands, the nine inputs and the separate UCI grade reference.']:
 doc.add_paragraph(x,style='List Bullet')
table(['Time','Activity','Evidence'],[['0:00 to 1:30','Problem and users','Retrospective performance bands for staff review'],['1:30 to 3:00','Design and data','Kaggle source, nine inputs and grouped split'],['3:00 to 4:30','Candidate comparison','Logistic regression wins grouped validation'],['4:30 to 7:00','Individual prediction','Name, observed activity and predicted band'],['7:00 to 10:00','Class dashboard','Historical label, filters, counts and CSV'],['10:00 to 12:00','Model quality','Accuracy, macro F1 and confusion matrix'],['12:00 to 13:30','GitHub, notebook and Drive','Revision, shared code and report ZIP'],['13:30 to 15:00','Conclusions and questions','92 tests, limitations and future data']],[1.2,2.05,3.5])
doc.add_page_break();doc.add_heading('Checks and discussion',1)
table(['Action','Expected result'],[['Sign in','Kaggle workspace is the default'],['Choose Sign up','Additional staff registration remains available'],['Inspect subjects','Twelve source-supported Kaggle topics'],['Predict an individual','Low, Middle or High band; no invented exact mark'],['Historical class view','96 test records before filters'],['Search name or ID','Filtered review table and matching CSV'],['Select UCI grade reference','Separate Mathematics and Portuguese grade models'],['Run notebook','Kaggle training, charts and exports, plus UCI reference'],['Inspect Git revision','Traceable application and experiment version']],[2.7,4.05])

doc.add_heading('Questions to prepare',2)
p('Why Kaggle xAPI? It has documented LMS origins, activity fields and twelve recorded topics. Why classify bands? Exact marks are absent. Why logistic regression? It has the highest grouped validation macro F1. Does 76% accuracy guarantee the next result? No. Are names predictors? No. Is it validated early warning? No, the source lacks cutoff dates and student identifiers. What does Drive save? Explicit project and report snapshots; sharing remains a separate permission.')
p('The existing presentation is the earlier UCI edition. Revise it for the Kaggle results before delivering this demo. The current report and guide contain the new evidence.')

doc.add_heading('Fallbacks and record of delivery',2)
p('If Streamlit fails, check the interpreter and requirements. If internet access fails, demonstrate bundled local subjects and prepared notebook outputs. If Google authorization fails, show the connection cell and identify the pending cloud step. Do not report a failed live step as successful.')
p('After presentation, record the actual date, participants, supervisor feedback and observed issues. Mark the assessed demo complete only after delivery. Update substantive claims when the implementation or evidence changes.')
doc.save(OUT/'StudentInsight_Demo_Guide.docx')
print('Created report and demo guide')
